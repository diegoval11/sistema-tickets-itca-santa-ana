from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth import login, logout
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.core.mail import send_mail
from django.conf import settings
from django.utils import timezone
from django.db.models import Q, Count
from datetime import timedelta, datetime
from .models import CustomUser, Ticket, TicketAttachment, TicketHistory, Notification
from .forms import AccessCodeLoginForm, UserCreationFormByTechnician, TicketCreationForm, TicketUpdateForm
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import csv


def login_view(request):
    if request.user.is_authenticated:
        if request.user.role == 'TECNICO':
            return redirect('dashboard')
        return redirect('my_tickets')
    
    if request.method == 'POST':
        form = AccessCodeLoginForm(request.POST)
        if form.is_valid():
            institutional_email = form.cleaned_data['institutional_email']
            access_code = form.cleaned_data['access_code']
            remember_me = form.cleaned_data['remember_me']
            
            try:
                user = CustomUser.objects.get(
                    institutional_email=institutional_email,
                    access_code=access_code,
                    is_code_active=True
                )
                
                login(request, user)
                
                if remember_me:
                    request.session.set_expiry(settings.SESSION_COOKIE_AGE)  # 30 dias
                else:
                    request.session.set_expiry(0) 
                
                messages.success(request, f'Bienvenido, {user.get_full_name()}!')
                
                if user.role == 'TECNICO':
                    return redirect('dashboard')
                return redirect('my_tickets')
                
            except CustomUser.DoesNotExist:
                form.add_error(None, 'Correo o código de acceso inválido.')
    else:
        form = AccessCodeLoginForm()
    
    return render(request, 'login.html', {'form': form})


def logout_view(request):
    logout(request)
    messages.info(request, 'Has cerrado sesión correctamente.')
    return redirect('login')


@login_required
def my_tickets(request):
    if request.user.role == 'TECNICO':
        return redirect('dashboard')
    
    status_filter = request.GET.get('status', '')
    search_query = request.GET.get('search', '')
    
    tickets = Ticket.objects.filter(user=request.user)
    
    if status_filter:
        tickets = tickets.filter(status=status_filter)
    
    if search_query:
        tickets = tickets.filter(
            Q(ticket_number__icontains=search_query) |
            Q(description__icontains=search_query)
        )
    
    unread_count = request.user.notifications.filter(is_read=False).count()
    
    context = {
        'tickets': tickets,
        'status_filter': status_filter,
        'search_query': search_query,
        'unread_count': unread_count,
    }
    
    return render(request, 'my_tickets.html', context)


@login_required
def create_ticket(request):
    if request.user.role == 'TECNICO':
        return redirect('dashboard')
    
    if request.method == 'POST':
        form = TicketCreationForm(request.POST, request.FILES)
        if form.is_valid():
            ticket = form.save(commit=False)
            ticket.user = request.user
            ticket.save()
            
            photos = request.FILES.getlist('photos')
            
            if len(photos) > settings.MAX_PHOTOS_PER_TICKET:
                messages.error(request, f'Máximo {settings.MAX_PHOTOS_PER_TICKET} fotos permitidas.')
                ticket.delete()
                return render(request, 'create_ticket.html', {
                    'form': form,
                    'unread_count': request.user.notifications.filter(is_read=False).count(),
                })
            
            for photo in photos:
                if photo.content_type not in settings.ALLOWED_IMAGE_TYPES:
                    messages.error(request, 'Solo se permiten archivos JPG y PNG.')
                    ticket.delete()
                    return render(request, 'create_ticket.html', {
                        'form': form,
                        'unread_count': request.user.notifications.filter(is_read=False).count(),
                    })
                
                if photo.size > settings.MAX_UPLOAD_SIZE:
                    messages.error(request, 'Cada foto debe ser menor a 5MB.')
                    ticket.delete()
                    return render(request, 'create_ticket.html', {
                        'form': form,
                        'unread_count': request.user.notifications.filter(is_read=False).count(),
                    })
                
                TicketAttachment.objects.create(ticket=ticket, image=photo)
            
            TicketHistory.objects.create(
                ticket=ticket,
                user=request.user,
                action='CREATED',
                comment='Ticket creado'
            )
            
            technicians = CustomUser.objects.filter(role='TECNICO')
            for tech in technicians:
                Notification.objects.create(
                    user=tech,
                    ticket=ticket,
                    title='Nuevo Ticket',
                    message=f'Nuevo ticket {ticket.ticket_number} creado por {request.user.get_full_name()}'
                )
                
                send_mail(
                    subject=f'Nuevo Ticket: {ticket.ticket_number}',
                    message=f'Se ha creado un nuevo ticket.\n\nNúmero: {ticket.ticket_number}\nUsuario: {request.user.get_full_name()}\nCategoría: {ticket.get_category_display()}\nPrioridad: {ticket.get_priority_display()}\n\nDescripción: {ticket.description}',
                    from_email=settings.DEFAULT_FROM_EMAIL,
                    recipient_list=[tech.institutional_email],
                    fail_silently=True,
                )
            
            messages.success(request, f'Ticket {ticket.ticket_number} creado exitosamente.')
            return redirect('ticket_detail', ticket_id=ticket.id)
    else:
        form = TicketCreationForm()
    
    unread_count = request.user.notifications.filter(is_read=False).count()
    
    return render(request, 'create_ticket.html', {
        'form': form,
        'unread_count': unread_count,
    })


@login_required
def ticket_detail(request, ticket_id):
    ticket = get_object_or_404(Ticket, id=ticket_id)
    
    if request.user.role != 'TECNICO' and ticket.user != request.user:
        messages.error(request, 'No tienes permiso para ver este ticket.')
        return redirect('my_tickets')
    
    history = ticket.history.all()
    
    attachments = ticket.attachments.all()
    
    unread_count = request.user.notifications.filter(is_read=False).count()
    
    context = {
        'ticket': ticket,
        'history': history,
        'attachments': attachments,
        'unread_count': unread_count,
    }
    
    return render(request, 'ticket_detail.html', context)


@login_required
def dashboard(request):
    if request.user.role != 'TECNICO':
        return redirect('my_tickets')
    
    tickets = Ticket.objects.all()
    
    status_filter = request.GET.get('status', '')
    priority_filter = request.GET.get('priority', '')
    category_filter = request.GET.get('category', '')
    search_query = request.GET.get('search', '')
    
    if status_filter:
        tickets = tickets.filter(status=status_filter)
    
    if priority_filter:
        tickets = tickets.filter(priority=priority_filter)
    
    if category_filter:
        tickets = tickets.filter(category=category_filter)
    
    if search_query:
        tickets = tickets.filter(
            Q(ticket_number__icontains=search_query) |
            Q(description__icontains=search_query) |
            Q(user__first_name__icontains=search_query) |
            Q(user__last_name__icontains=search_query)
        )
    
    total_tickets = Ticket.objects.count()
    open_tickets = Ticket.objects.filter(status='ABIERTO').count()
    in_progress_tickets = Ticket.objects.filter(status='EN_PROGRESO').count()
    closed_tickets = Ticket.objects.filter(status='CERRADO').count()
    rejected_tickets = Ticket.objects.filter(status='RECHAZADO').count()
    archived_tickets = Ticket.objects.filter(status='ARCHIVADO').count()
    
    tickets_by_category = Ticket.objects.values('category').annotate(count=Count('id'))
    
    tickets_by_priority = Ticket.objects.values('priority').annotate(count=Count('id'))
    
    closed_with_time = Ticket.objects.filter(status='CERRADO', closed_at__isnull=False)
    if closed_with_time.exists():
        total_time = sum([(t.closed_at - t.created_at).total_seconds() for t in closed_with_time])
        avg_resolution_hours = (total_time / closed_with_time.count()) / 3600
    else:
        avg_resolution_hours = 0
    
    now = timezone.now()
    month_start = now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    tickets_this_month = Ticket.objects.filter(created_at__gte=month_start).count()
    
    quarter_start = now.replace(month=((now.month - 1) // 3) * 3 + 1, day=1, hour=0, minute=0, second=0, microsecond=0)
    tickets_this_quarter = Ticket.objects.filter(created_at__gte=quarter_start).count()
    
    unread_count = request.user.notifications.filter(is_read=False).count()
    
    context = {
        'tickets': tickets[:20],  # Show latest 20
        'total_tickets': total_tickets,
        'open_tickets': open_tickets,
        'in_progress_tickets': in_progress_tickets,
        'closed_tickets': closed_tickets,
        'rejected_tickets': rejected_tickets,
        'archived_tickets': archived_tickets,
        'tickets_by_category': tickets_by_category,
        'tickets_by_priority': tickets_by_priority,
        'avg_resolution_hours': round(avg_resolution_hours, 1),
        'tickets_this_month': tickets_this_month,
        'tickets_this_quarter': tickets_this_quarter,
        'status_filter': status_filter,
        'priority_filter': priority_filter,
        'category_filter': category_filter,
        'search_query': search_query,
        'unread_count': unread_count,
    }
    
    return render(request, 'dashboard.html', context)


@login_required
def user_management(request):
    """View for technicians to manage users"""
    if request.user.role != 'TECNICO':
        return redirect('my_tickets')
    
    users = CustomUser.objects.filter(role='USUARIO').order_by('-date_joined')
    
    unread_count = request.user.notifications.filter(is_read=False).count()
    
    return render(request, 'user_management.html', {
        'users': users,
        'unread_count': unread_count,
    })


@login_required
def create_user(request):
    """View for technicians to create new users"""
    if request.user.role != 'TECNICO':
        return redirect('my_tickets')
    
    if request.method == 'POST':
        form = UserCreationFormByTechnician(request.POST)
        if form.is_valid():
            user = form.save(commit=False)
            user.role = 'USUARIO'
            user.set_unusable_password()  # No password needed
            user.save()
            
            # Send access code via email
            send_mail(
                subject='Código de Acceso - Sistema de Tickets ITCA FEPADE',
                message=f'Hola {user.get_full_name()},\n\nTu cuenta ha sido creada en el Sistema de Tickets de ITCA FEPADE.\n\nTu código de acceso es: {user.access_code}\n\nPuedes iniciar sesión en: http://localhost:8000\n\nCorreo: {user.institutional_email}\nCódigo: {user.access_code}\n\nSaludos,\nEquipo de Soporte Técnico',
                from_email=settings.DEFAULT_FROM_EMAIL,
                recipient_list=[user.institutional_email],
                fail_silently=True,
            )
            
            user.code_sent_count = 1
            user.save()
            
            messages.success(request, f'Usuario {user.get_full_name()} creado. Código enviado a {user.institutional_email}.')
            return redirect('user_management')
    else:
        form = UserCreationFormByTechnician()
    
    unread_count = request.user.notifications.filter(is_read=False).count()
    
    return render(request, 'create_user.html', {
        'form': form,
        'unread_count': unread_count,
    })


@login_required
def resend_code(request, user_id):
    """View for technicians to resend access code"""
    if request.user.role != 'TECNICO':
        return redirect('my_tickets')
    
    user = get_object_or_404(CustomUser, id=user_id)
    
    if user.code_sent_count >= 3:
        messages.warning(request, f'El código ya ha sido enviado 3 veces. Considera regenerar el código.')
    else:
        # Send email
        send_mail(
            subject='Código de Acceso - Sistema de Tickets ITCA FEPADE',
            message=f'Hola {user.get_full_name()},\n\nTu código de acceso es: {user.access_code}\n\nPuedes iniciar sesión en: http://localhost:8000\n\nCorreo: {user.institutional_email}\nCódigo: {user.access_code}\n\nSaludos,\nEquipo de Soporte Técnico',
            from_email=settings.DEFAULT_FROM_EMAIL,
            recipient_list=[user.institutional_email],
            fail_silently=True,
        )
        
        user.code_sent_count += 1
        user.save()
        
        messages.success(request, f'Código reenviado a {user.institutional_email}.')
    
    return redirect('user_management')


@login_required
def regenerate_code(request, user_id):
    """View for technicians to regenerate access code"""
    if request.user.role != 'TECNICO':
        return redirect('my_tickets')
    
    user = get_object_or_404(CustomUser, id=user_id)
    
    # Regenerate code
    new_code = user.regenerate_code()
    
    # Send email with new code
    send_mail(
        subject='Nuevo Código de Acceso - Sistema de Tickets ITCA FEPADE',
        message=f'Hola {user.get_full_name()},\n\nTu código de acceso ha sido regenerado.\n\nTu nuevo código es: {new_code}\n\nPuedes iniciar sesión en: http://localhost:8000\n\nCorreo: {user.institutional_email}\nCódigo: {new_code}\n\nSaludos,\nEquipo de Soporte Técnico',
        from_email=settings.DEFAULT_FROM_EMAIL,
        recipient_list=[user.institutional_email],
        fail_silently=True,
    )
    
    user.code_sent_count = 1
    user.save()
    
    messages.success(request, f'Código regenerado y enviado a {user.institutional_email}.')
    
    return redirect('user_management')


@login_required
def update_ticket(request, ticket_id):
    """View for technicians to update ticket status"""
    if request.user.role != 'TECNICO':
        return redirect('my_tickets')
    
    ticket = get_object_or_404(Ticket, id=ticket_id)
    
    if request.method == 'POST':
        form = TicketUpdateForm(request.POST, instance=ticket)
        if form.is_valid():
            old_status = ticket.status
            ticket = form.save()
            
            # Create history entry based on status change
            if old_status != ticket.status:
                if ticket.status == 'EN_PROGRESO':
                    action = 'STATUS_CHANGED'
                    comment = 'Ticket aceptado y en progreso'
                    
                    # If visit scheduled
                    if ticket.visit_date and ticket.visit_time:
                        action = 'VISIT_SCHEDULED'
                        comment = f'Visita programada para {ticket.visit_date} a las {ticket.visit_time}'
                        
                        # Notify user
                        Notification.objects.create(
                            user=ticket.user,
                            ticket=ticket,
                            title='Visita Programada',
                            message=f'Tu ticket {ticket.ticket_number} tiene una visita programada para {ticket.visit_date} a las {ticket.visit_time}'
                        )
                        
                        # Send email
                        send_mail(
                            subject=f'Visita Programada - Ticket {ticket.ticket_number}',
                            message=f'Hola {ticket.user.get_full_name()},\n\nTu ticket {ticket.ticket_number} ha sido aceptado.\n\nVisita programada: {ticket.visit_date} a las {ticket.visit_time}\n\nSaludos,\nEquipo de Soporte Técnico',
                            from_email=settings.DEFAULT_FROM_EMAIL,
                            recipient_list=[ticket.user.institutional_email],
                            fail_silently=True,
                        )
                
                elif ticket.status == 'RECHAZADO':
                    action = 'REJECTED'
                    comment = f'Ticket rechazado: {ticket.rejection_reason}'
                    
                    Notification.objects.create(
                        user=ticket.user,
                        ticket=ticket,
                        title='Ticket Rechazado',
                        message=f'Tu ticket {ticket.ticket_number} ha sido rechazado. Motivo: {ticket.rejection_reason}'
                    )
                    
                    send_mail(
                        subject=f'Ticket Rechazado - {ticket.ticket_number}',
                        message=f'Hola {ticket.user.get_full_name()},\n\nTu ticket {ticket.ticket_number} ha sido rechazado.\n\nMotivo: {ticket.rejection_reason}\n\nSaludos,\nEquipo de Soporte Técnico',
                        from_email=settings.DEFAULT_FROM_EMAIL,
                        recipient_list=[ticket.user.institutional_email],
                        fail_silently=True,
                    )
                
                elif ticket.status == 'CERRADO':
                    action = 'CLOSED'
                    comment = f'Ticket cerrado: {ticket.closure_note or "Sin nota"}'
                    ticket.closed_at = timezone.now()
                    ticket.save()
                    
                    Notification.objects.create(
                        user=ticket.user,
                        ticket=ticket,
                        title='Ticket Cerrado',
                        message=f'Tu ticket {ticket.ticket_number} ha sido cerrado. {ticket.closure_note or ""}'
                    )
                    
                    send_mail(
                        subject=f'Ticket Cerrado - {ticket.ticket_number}',
                        message=f'Hola {ticket.user.get_full_name()},\n\nTu ticket {ticket.ticket_number} ha sido cerrado.\n\n{ticket.closure_note or ""}\n\nSaludos,\nEquipo de Soporte Técnico',
                        from_email=settings.DEFAULT_FROM_EMAIL,
                        recipient_list=[ticket.user.institutional_email],
                        fail_silently=True,
                    )
                
                else:
                    action = 'STATUS_CHANGED'
                    comment = f'Estado cambiado a {ticket.get_status_display()}'
                
                TicketHistory.objects.create(
                    ticket=ticket,
                    user=request.user,
                    action=action,
                    comment=comment
                )
            
            messages.success(request, 'Ticket actualizado exitosamente.')
            return redirect('ticket_detail', ticket_id=ticket.id)
    else:
        form = TicketUpdateForm(instance=ticket)
    
    unread_count = request.user.notifications.filter(is_read=False).count()
    
    return render(request, 'update_ticket.html', {
        'form': form,
        'ticket': ticket,
        'unread_count': unread_count,
    })


@login_required
def notifications(request):
    user_notifications = request.user.notifications.all()
    
    unread_count = request.user.notifications.filter(is_read=False).count()
    
    return render(request, 'notifications.html', {
        'notifications': user_notifications,
        'unread_count': unread_count,
    })


@login_required
def mark_notification_read(request, notification_id):
    notification = get_object_or_404(Notification, id=notification_id, user=request.user)
    notification.is_read = True
    notification.save()
    
    return redirect('notifications')


@login_required
def export_tickets(request):
    if request.user.role != 'TECNICO':
        return redirect('my_tickets')
    
    unread_count = request.user.notifications.filter(is_read=False).count()
    
    if request.method == 'GET':
        return render(request, 'export_tickets.html', {
            'unread_count': unread_count,
        })
    
    start_date = request.POST.get('start_date')
    end_date = request.POST.get('end_date')
    export_format = request.POST.get('format', 'xlsx')
    
    # Build queryset
    tickets = Ticket.objects.all()
    
    if start_date:
        tickets = tickets.filter(created_at__gte=start_date)
    
    if end_date:
        tickets = tickets.filter(created_at__lte=end_date)
    
    if export_format == 'csv':
        return export_to_csv(tickets, start_date, end_date)
    else:
        return export_to_excel(tickets, start_date, end_date)


def export_to_csv(tickets, start_date, end_date):
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = f'attachment; filename="tickets_{datetime.now().strftime("%Y%m%d")}.csv"'
    
    writer = csv.writer(response)
    
    writer.writerow([
        'Número de Ticket',
        'Usuario',
        'Correo',
        'Categoría',
        'Prioridad',
        'Estado',
        'Equipo Afectado',
        'Número de Serie',
        'Descripción',
        'Fecha de Creación',
        'Fecha de Cierre',
        'Fecha de Visita',
        'Hora de Visita',
        'Motivo de Rechazo',
        'Nota de Cierre'
    ])
    
    # Write data
    for ticket in tickets:
        writer.writerow([
            ticket.ticket_number,
            ticket.user.get_full_name(),
            ticket.user.institutional_email,
            ticket.get_category_display(),
            ticket.get_priority_display(),
            ticket.get_status_display(),
            ticket.affected_equipment,
            ticket.serial_number or '',
            ticket.description,
            ticket.created_at.strftime('%d/%m/%Y %H:%M'),
            ticket.closed_at.strftime('%d/%m/%Y %H:%M') if ticket.closed_at else '',
            ticket.visit_date.strftime('%d/%m/%Y') if ticket.visit_date else '',
            ticket.visit_time.strftime('%H:%M') if ticket.visit_time else '',
            ticket.rejection_reason or '',
            ticket.closure_note or ''
        ])
    
    return response


def export_to_excel(tickets, start_date, end_date):
    wb = Workbook()
    
    ws_summary = wb.active
    ws_summary.title = "Resumen"
    
    header_fill = PatternFill(start_color="D4A574", end_color="D4A574", fill_type="solid")
    header_font = Font(bold=True, color="000000")
    
    ws_summary['A1'] = 'REPORTE DE TICKETS - ITCA FEPADE'
    ws_summary['A1'].font = Font(bold=True, size=16)
    ws_summary.merge_cells('A1:D1')
    
    ws_summary['A2'] = f'Período: {start_date or "Inicio"} - {end_date or "Hoy"}'
    ws_summary['A2'].font = Font(italic=True)
    ws_summary.merge_cells('A2:D2')
    
    ws_summary['A4'] = 'Métrica'
    ws_summary['B4'] = 'Valor'
    ws_summary['A4'].fill = header_fill
    ws_summary['B4'].fill = header_fill
    ws_summary['A4'].font = header_font
    ws_summary['B4'].font = header_font
    
    metrics = [
        ('Total de Tickets', tickets.count()),
        ('Abiertos', tickets.filter(status='ABIERTO').count()),
        ('En Progreso', tickets.filter(status='EN_PROGRESO').count()),
        ('Rechazados', tickets.filter(status='RECHAZADO').count()),
        ('Cerrados', tickets.filter(status='CERRADO').count()),
        ('Archivados', tickets.filter(status='ARCHIVADO').count()),
    ]
    
    row = 5
    for metric, value in metrics:
        ws_summary[f'A{row}'] = metric
        ws_summary[f'B{row}'] = value
        row += 1
    
    # Calculate average resolution time
    closed_tickets = tickets.filter(status='CERRADO', closed_at__isnull=False)
    if closed_tickets.exists():
        total_time = sum([(t.closed_at - t.created_at).total_seconds() for t in closed_tickets])
        avg_hours = (total_time / closed_tickets.count()) / 3600
        ws_summary[f'A{row}'] = 'Tiempo Promedio de Resolución (horas)'
        ws_summary[f'B{row}'] = round(avg_hours, 1)
    
    ws_summary.column_dimensions['A'].width = 40
    ws_summary.column_dimensions['B'].width = 15
    
    ws_tickets = wb.create_sheet(title="Tickets")
    
    headers = [
        'Número', 'Usuario', 'Correo', 'Categoría', 'Prioridad', 'Estado',
        'Equipo', 'Serie', 'Descripción', 'Creado', 'Cerrado',
        'Visita Fecha', 'Visita Hora', 'Motivo Rechazo', 'Nota Cierre'
    ]
    
    for col, header in enumerate(headers, start=1):
        cell = ws_tickets.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
    
    # Data
    for row, ticket in enumerate(tickets, start=2):
        ws_tickets.cell(row=row, column=1, value=ticket.ticket_number)
        ws_tickets.cell(row=row, column=2, value=ticket.user.get_full_name())
        ws_tickets.cell(row=row, column=3, value=ticket.user.institutional_email)
        ws_tickets.cell(row=row, column=4, value=ticket.get_category_display())
        ws_tickets.cell(row=row, column=5, value=ticket.get_priority_display())
        ws_tickets.cell(row=row, column=6, value=ticket.get_status_display())
        ws_tickets.cell(row=row, column=7, value=ticket.affected_equipment)
        ws_tickets.cell(row=row, column=8, value=ticket.serial_number or '')
        ws_tickets.cell(row=row, column=9, value=ticket.description)
        ws_tickets.cell(row=row, column=10, value=ticket.created_at.strftime('%d/%m/%Y %H:%M'))
        ws_tickets.cell(row=row, column=11, value=ticket.closed_at.strftime('%d/%m/%Y %H:%M') if ticket.closed_at else '')
        ws_tickets.cell(row=row, column=12, value=ticket.visit_date.strftime('%d/%m/%Y') if ticket.visit_date else '')
        ws_tickets.cell(row=row, column=13, value=ticket.visit_time.strftime('%H:%M') if ticket.visit_time else '')
        ws_tickets.cell(row=row, column=14, value=ticket.rejection_reason or '')
        ws_tickets.cell(row=row, column=15, value=ticket.closure_note or '')
    
    # Adjust column widths
    for col in range(1, 16):
        ws_tickets.column_dimensions[chr(64 + col)].width = 15
    ws_tickets.column_dimensions['I'].width = 40  # Description
    
    # Save to response
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="tickets_{datetime.now().strftime("%Y%m%d")}.xlsx"'
    wb.save(response)
    
    return response
